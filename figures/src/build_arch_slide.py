"""Before/after architecture slide: TimeXer -> TimeXer-DINOv3 (style follows the TimeXer paper figure).
Usage: python figures/src/build_arch_slide.py <img_dir> <out.pptx>   (needs python-pptx; images from make_imgs.py)"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from lxml import etree

IMG, OUT = sys.argv[1], sys.argv[2]
FONT_B, FONT = "Noto Sans Bold", "Noto Sans"
C = dict(embed="FFC2E8", ln="D9F4CD", ffn="C2CDFF", attn="FFC099", proj="A6A6A6", panel="F2F1F1",
         dino="B6C3FF", adapt="E8E8E8", pool="FFDE59", red="FF3131", call="FFD21F", blue="5170FF",
         muted="595959", white="FFFFFF")
RADIUS = 0.12  # corner radius in inches, same for every block

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(18288000), Emu(10287000)   # 20 x 11.25 in, same as the template
sh = prs.slides.add_slide(prs.slide_layouts[6]).shapes


def rgb(h): return RGBColor.from_string(h)


def _effects(shape, shadow):
    """Replace the shape's effects with nothing or a soft bottom-right drop shadow."""
    spPr = shape._element.spPr
    for e in spPr.findall(qn("a:effectLst")):
        spPr.remove(e)
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    if shadow:
        s = etree.SubElement(eff, qn("a:outerShdw"), blurRad="63500", dist="50800", dir="2700000",
                             algn="tl", rotWithShape="0")
        etree.SubElement(etree.SubElement(s, qn("a:srgbClr"), val="000000"), qn("a:alpha"), val="35000")


def _fill_alpha(shape, alpha_pct):
    clr = shape._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    etree.SubElement(clr, qn("a:alpha"), val=str(int(alpha_pct * 1000)))


def _runs(tf, lines, align):
    for k, (t, size, bold, color, italic) in enumerate(lines):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = t
        r.font.size = Pt(size); r.font.name = FONT_B if bold else FONT; r.font.bold = bold
        r.font.italic = italic; r.font.color.rgb = rgb(color)


def text(x, y, w, h, t, size=18, bold=True, color="000000", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, italic=False):
    tb = sh.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    _runs(tf, [(line, size, bold, color, italic) for line in t.split("\n")], align)
    return tb


def block(x, y, w, h, label, fill, size=18, sub=None, line="000000", lw=2.0, dash=None, shadow=True,
          shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=RADIUS):
    b = sh.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        b.adjustments[0] = min(0.5, radius / min(w, h))
    if fill:
        b.fill.solid(); b.fill.fore_color.rgb = rgb(fill)
    else:
        b.fill.background()
    if line:
        b.line.color.rgb = rgb(line); b.line.width = Pt(lw)
        if dash:
            b.line.dash_style = dash
    else:
        b.line.fill.background()
    _effects(b, shadow)
    if label is not None:
        tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, Inches(0.04))
        lines = [(label, size, True, "000000", False)]
        if sub:
            lines += [(s, size - 6, False, "000000", False) for s in sub.split("\n")]
        _runs(tf, lines, PP_ALIGN.CENTER)
    return b


def arrow(x1, y1, x2, y2, color="000000", w=2.0, dash=None, head=True):
    c = sh.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(w)
    if dash:
        c.line.dash_style = dash
    if head:
        te = etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd"))
        te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")
    return c


def new_tag(x, y):
    t = block(x, y, 0.62, 0.3, "NEW", C["red"], size=11, line=C["red"], lw=1, shadow=False, radius=0.06)
    t.text_frame.paragraphs[0].runs[0].font.color.rgb = rgb(C["white"])


def x_badge(cx, cy, r=0.27):
    """Red circle with a white cross: 'removed'."""
    block(cx - r, cy - r, 2 * r, 2 * r, None, C["red"], line=C["white"], lw=2, shape=MSO_SHAPE.OVAL)
    d = r * 0.45
    arrow(cx - d, cy - d, cx + d, cy + d, color=C["white"], w=3.5, head=False)
    arrow(cx - d, cy + d, cx + d, cy - d, color=C["white"], w=3.5, head=False)


def pic(name, x, y, w=None, h=None):
    return sh.add_picture(f"{IMG}/{name}.png", Inches(x), Inches(y),
                          width=Inches(w) if w else None, height=Inches(h) if h else None)


bh = 0.55  # block height

# ---------------- (a) TimeXer ----------------
text(0.4, 0.3, 7.3, 0.55, "(a) TimeXer (baseline)", size=26)
ax, aw = 3.3, 3.4
cx = ax + aw / 2
text(cx - 1, 1.05, 2, 0.4, "Output", size=18)
arrow(cx, 1.75, cx, 1.42)
block(ax + 0.35, 1.75, aw - 0.7, bh, "Projection", C["proj"])
block(ax - 0.25, 2.75, aw + 0.5, 6.0, None, C["panel"], line="000000", lw=1.25, radius=0.08)
text(ax + aw + 0.4, 2.75, 0.8, 0.45, "L ×", size=20)
stack = [("LayerNorm", C["ln"]), ("Feed-Forward", C["ffn"]), ("LayerNorm", C["ln"]),
         ("Cross-Attention", C["attn"]), ("LayerNorm", C["ln"]), ("Self-Attention", C["attn"])]
ys = [2.95 + 1.0 * k for k in range(len(stack))]
arrow(cx, 2.95, cx, 2.3)
for (lab, col), y in zip(stack, ys):
    block(ax + 0.1, y, aw - 0.2, bh, lab, col, shadow=False)
for y in ys[1:]:
    arrow(cx, y, cx, y - 0.45)
block(ax + 0.35, 9.2, aw - 0.7, bh, "Embedding", C["embed"], shadow=False)
arrow(cx, 9.2, cx, ys[-1] + bh)
text(cx - 1.6, 10.1, 3.2, 0.36, "Endogenous Input", size=18)
arrow(cx, 10.12, cx, 9.75)
text(cx - 1.6, 10.46, 3.2, 0.3, "patches + global token", size=12, bold=False, color=C["muted"])
ye = ys[3]
block(0.45, ye, 1.95, bh, "Embedding", C["embed"])
arrow(2.4, ye + bh / 2, ax + 0.1, ye + bh / 2)
text(0.25, ye + 0.85, 2.35, 0.36, "Exogenous Input", size=16)
arrow(1.42, ye + 0.87, 1.42, ye + bh)
pic("endo", 0.5, 8.9, w=2.2)
text(0.4, 9.65, 2.4, 0.3, "target, 168 h", size=12, bold=False, color=C["muted"])
# replaced region: fade it, outline it, cross it out, point to its replacement
rx0, rx1, ry0, ry1 = ax - 0.35, ax + aw + 0.35, 2.62, 9.88
fade = block(rx0, ry0, rx1 - rx0, ry1 - ry0, None, C["white"], line=None, shadow=False, radius=0.1)
_fill_alpha(fade, 45)
block(rx0, ry0, rx1 - rx0, ry1 - ry0, None, None, line=C["red"], lw=1.75, dash=MSO_LINE.DASH, shadow=False, radius=0.1)
x_badge(rx1, ry0)
text(rx0, ry0 - 0.42, 2.6, 0.34, "removed", size=15, color=C["red"], align=PP_ALIGN.LEFT)

arrow(8.05, 0.95, 8.05, 10.55, color="BFBFBF", w=1.5, dash=MSO_LINE.DASH, head=False)

# ---------------- (b) TimeXer-DINOv3 ----------------
text(8.4, 0.3, 11.2, 0.55, "(b) TimeXer-DINOv3 (ours)", size=26, align=PP_ALIGN.LEFT)
bx, bw = 10.75, 3.9
bcx = bx + bw / 2
text(bcx - 1, 1.05, 2, 0.4, "Output", size=18)
arrow(bcx, 1.75, bcx, 1.42)
block(bx + 0.45, 1.75, bw - 0.9, bh, "Projection", C["proj"])
block(bx - 0.25, 2.55, bw + 0.5, 4.1, None, C["panel"], line="000000", lw=1.25, radius=0.08)
rows = [(2.75, "Pool 121 → 7 tokens", C["pool"], None),
        (3.75, "Self-Attention × 1", C["attn"], "121 patches + fused CLS"),
        (4.75, "Cross-Attention × 1", C["attn"], "query = CLS"),
        (5.75, "Linear Adapter", C["adapt"], "384 → 256")]
for y, lab, col, sub in rows:
    block(bx, y, bw, bh + (0.12 if sub else 0), lab, col, sub=sub, shadow=False)
    new_tag(bx + bw - 0.35, y - 0.17)
arrow(bcx, 2.75, bcx, 2.3)
for yt, yb in [(3.75, 3.3), (4.75, 4.42), (5.75, 5.42)]:
    arrow(bcx, yt, bcx, yb)
# patch tokens bypass the CLS cross-attention
ex = bx + bw + 0.48
arrow(bx + bw, 6.08, ex, 6.08, color=C["blue"], w=1.75, head=False)
arrow(ex, 6.08, ex, 4.09, color=C["blue"], w=1.75, head=False)
arrow(ex, 4.09, bx + bw, 4.09, color=C["blue"], w=1.75)
pt = text(ex + 0.04, 4.45, 0.3, 1.3, "patches", size=12, bold=False, color=C["blue"])
pt.text_frame._txBody.bodyPr.set("vert", "vert270")
text(bcx - 0.75, 5.4, 0.65, 0.32, "CLS", size=12, bold=False, color=C["blue"], align=PP_ALIGN.RIGHT)
dy, dh = 6.95, 1.6
block(bx - 0.2, dy, bw + 0.4, dh, "DINOv3 ViT-S/16", C["dino"], size=20,
      sub="12 blocks · pretrained on LVD-1689M\nR1: frozen · R2: top-2 blocks tuned")
new_tag(bx + bw - 0.15, dy - 0.17)
arrow(bcx, dy, bcx, 6.42)
block(bx + 0.25, 9.2, bw - 0.5, bh, "TS2Image", C["embed"])
new_tag(bx + bw - 0.55, 9.03)
arrow(bcx, 9.2, bcx, dy + dh)
text(bcx - 1.6, 10.1, 3.2, 0.36, "Endogenous Input", size=18)
arrow(bcx, 10.12, bcx, 9.75)
text(bcx - 1.6, 10.46, 3.2, 0.3, "RevIN-normalised window", size=12, bold=False, color=C["muted"])
ye = 4.75
block(8.4, ye, 1.95, bh, "Embedding", C["embed"])
arrow(10.35, ye + bh / 2 + 0.06, bx, ye + bh / 2 + 0.06)
text(8.2, ye + 0.85, 2.35, 0.36, "Exogenous Input", size=16)
arrow(9.37, ye + 0.87, 9.37, ye + bh)
pic("exo1", 8.55, ye + 1.35, w=1.7)
pic("exo2", 8.55, ye + 1.9, w=1.7)
text(8.3, ye + 2.5, 2.2, 0.3, "load & wind forecasts", size=12, bold=False, color=C["muted"])
text(8.3, ye - 0.42, 2.2, 0.35, "unchanged", size=13, bold=False, color=C["muted"], italic=True)

# the removed encoder in (a) is replaced by TS2Image + DINOv3 in (b)
ya = dy + dh / 2 + 0.25
arrow(rx1, ya, bx - 0.2, ya, color=C["red"], w=2.75)
text(rx1 + 0.1, ya - 0.4, bx - 0.3 - rx1 - 0.1, 0.34, "replaced by", size=14, color=C["red"])

# (c) 1D -> 2D imaging
ix, iw = 15.55, 4.05
block(ix, 8.05, iw, 2.55, None, C["white"], radius=0.08)
text(ix + 0.15, 8.12, iw - 0.3, 0.4, "(c) 1D → 2D imaging", size=16, align=PP_ALIGN.LEFT)
for k, (nm, lab) in enumerate([("gasf", "GASF"), ("gadf", "GADF"), ("fold", "Periodic fold")]):
    px = ix + 0.17 + k * 1.3
    pic(nm, px, 8.6, w=1.12, h=1.12)
    text(px - 0.15, 9.8, 1.42, 0.3, lab, size=12)
text(ix + 0.1, 10.15, iw - 0.2, 0.35, "168 h → 3 × 176 × 176 image", size=12, bold=False, color=C["muted"])
arrow(bx + bw - 0.25, 9.47, ix, 9.2, color=C["call"], w=2, dash=MSO_LINE.DASH, head=False)

# (d) backbone output tokens
block(ix, 4.55, iw, 3.15, None, C["white"], radius=0.08)
text(ix + 0.15, 4.62, iw - 0.3, 0.4, "(d) Backbone output tokens", size=16, align=PP_ALIGN.LEFT)
pic("patches", ix + 0.2, 5.15, w=1.9, h=1.9)
text(ix + 0.1, 7.1, 2.1, 0.3, "121 patches (11 × 11)", size=11, bold=False, color=C["muted"])
tx = ix + 2.3
block(tx, 5.2, 0.5, 0.5, None, "8FA3FF", lw=1.5, shadow=False, radius=0.05)
text(tx + 0.6, 5.2, 1.1, 0.5, "CLS: kept", size=12, align=PP_ALIGN.LEFT)
for k in range(4):
    block(tx + k * 0.38, 6.0, 0.32, 0.32, None, "D9D9D9", lw=1.2, shadow=False, radius=0.04)
arrow(tx - 0.05, 5.95, tx + 1.52, 6.37, color=C["red"], w=2.5, head=False)
arrow(tx - 0.05, 6.37, tx + 1.52, 5.95, color=C["red"], w=2.5, head=False)
text(tx - 0.05, 6.45, 1.7, 0.5, "4 registers:\ndropped", size=12, align=PP_ALIGN.LEFT, color=C["red"])
arrow(bx + bw + 0.2, 7.45, ix, 7.1, color=C["call"], w=2, dash=MSO_LINE.DASH, head=False)

# (e) summary of the change
block(ix, 1.75, iw, 2.3, None, C["white"], radius=0.08)
text(ix + 0.15, 1.82, iw - 0.3, 0.4, "(e) Change vs. (a)", size=16, align=PP_ALIGN.LEFT)
text(ix + 0.2, 2.3, iw - 0.35, 1.7,
     "Patch embedding + L encoder layers\n→ image + pretrained ViT\n\n"
     "Cross-attention in every layer\n→ once, after the backbone\n\n"
     "Loss, data, RevIN, exo embedding,\nprojection: unchanged",
     size=12, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)

# legend
ly = 10.86
new_tag(0.45, ly)
text(1.15, ly - 0.03, 2.6, 0.36, "new or replaced block", size=13, bold=False, align=PP_ALIGN.LEFT)
x_badge(4.05, ly + 0.15, r=0.17)
text(4.35, ly - 0.03, 3.2, 0.36, "removed from TimeXer", size=13, bold=False, align=PP_ALIGN.LEFT)
arrow(7.0, ly + 0.15, 7.7, ly + 0.15, color=C["call"], w=2, dash=MSO_LINE.DASH, head=False)
text(7.8, ly - 0.03, 2.5, 0.36, "zoom-in detail", size=13, bold=False, align=PP_ALIGN.LEFT)

prs.save(OUT)
print("saved", OUT)
